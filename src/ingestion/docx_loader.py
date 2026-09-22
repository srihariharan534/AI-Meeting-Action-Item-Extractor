"""DOCX document ingestion loader."""

import io
from typing import Union
from pathlib import Path
import docx


def load_docx(file_input: Union[str, Path, bytes]) -> str:
    """Extract raw text from DOCX document."""
    if isinstance(file_input, (str, Path)):
        path = Path(file_input)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {path}")
        doc = docx.Document(str(path))
    elif isinstance(file_input, bytes):
        doc = docx.Document(io.BytesIO(file_input))
    else:
        raise TypeError("Invalid file_input type for load_docx; must be str, Path, or bytes.")

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)
